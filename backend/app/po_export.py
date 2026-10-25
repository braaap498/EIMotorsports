"""
Purchase order files: a PDF laid out like the QuickBooks purchase order (company block top left,
Date / P.O. No. box top right, Vendor and Ship To boxes, then Item / Description / Qty / Rate / Amount
with the Total at the bottom), and an Excel file with the same information.

Needs two packages in requirements.txt:   reportlab    openpyxl

build_pdf(po) and build_xlsx(po) both take the same dict (see _example in the tests) and return bytes:
    po_number, date (a date), vendor_name, vendor_address, ship_to, note,
    lines: [{item, description, quantity, rate}]
"""
import io
from datetime import date
from decimal import ROUND_HALF_UP, Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib.utils import simpleSplit
from reportlab.pdfgen import canvas

# The business's own address, printed top left. The Ship To box defaults to the warehouse (it can be
# changed on each purchase order when it is created).
COMPANY_LINES = ["E & I Motorsports LLC", "7236 Valley Rd.", "Honor, MI 49640"]
DEFAULT_SHIP_TO = "E & I Motorsports LLC\n5572 Betsie River Rd\nInterlochen, MI 49643"


def _money(v) -> Decimal:
    return Decimal(str(v)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def line_amount(line: dict) -> Decimal:
    return _money(_money(line["rate"]) * int(line["quantity"]))


def po_total(lines: list[dict]) -> Decimal:
    return sum((line_amount(l) for l in lines), Decimal("0.00"))


def _date_text(d) -> str:
    d = d if isinstance(d, date) else date.today()
    return f"{d.month:02d}/{d.day:02d}/{d.year % 100:02d}"


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
_S = 72 / 70  # layout below was measured on a 70-dpi picture of the QuickBooks PO; this makes it points


def build_pdf(po: dict) -> bytes:
    buf = io.BytesIO()
    W, H = 612, 792
    c = canvas.Canvas(buf, pagesize=(W, H))
    c.setTitle(f"Purchase Order {po['po_number']}")

    def X(v):
        return v * _S

    def Y(v):  # distance from the TOP of the page, in layout units
        return H - v * _S

    def box(x1, y1, x2, y2, fill=None):
        if fill is not None:
            c.setFillGray(fill)
            c.rect(X(x1), Y(y2), X(x2) - X(x1), Y(y1) - Y(y2), stroke=0, fill=1)
            c.setFillGray(0)
        c.rect(X(x1), Y(y2), X(x2) - X(x1), Y(y1) - Y(y2), stroke=1, fill=0)

    def hline(x1, x2, y):
        c.line(X(x1), Y(y), X(x2), Y(y))

    def vline(x, y1, y2):
        c.line(X(x), Y(y1), X(x), Y(y2))

    c.setLineWidth(0.8)

    # ---- column layout of the item table
    TX1, TX2 = 36, 560
    COLS = [36, 113, 323, 402, 481, 560]
    BODY_BOTTOM, TOTAL_BOTTOM = 695, 727
    lines = po["lines"]
    total = po_total(lines)

    def page_one_header():
        c.setFont("Times-Roman", 10.5)
        for n, t in enumerate(COMPANY_LINES):
            c.drawString(X(38), Y(48 + (0 if n == 0 else 28 + (n - 1) * 16)), t)
        c.setFont("Helvetica-Bold", 18.5)
        c.drawString(X(404), Y(57), "Purchase Order")
        # date / PO number box
        box(428, 64, 560, 108)
        hline(428, 560, 88)
        vline(493, 64, 108)
        c.setFont("Helvetica", 8.5)
        c.drawCentredString(X(460), Y(80), "Date")
        c.drawCentredString(X(526), Y(80), "P.O. No.")
        c.setFont("Times-Roman", 8.5)
        c.drawCentredString(X(460), Y(101), _date_text(po.get("date")))
        c.drawCentredString(X(526), Y(101), str(po["po_number"]))
        # vendor and ship to
        for x1, x2, title, body in (
            (61, 301, "Vendor", f"{po['vendor_name']}\n{po.get('vendor_address') or ''}"),
            (319, 559, "Ship To", po.get("ship_to") or ""),
        ):
            box(x1, 126, x2, 231)
            hline(x1, x2, 147)
            c.setFont("Helvetica", 8.5)
            c.drawString(X(x1 + 12), Y(140), title)
            c.setFont("Times-Roman", 8)
            yy = 160
            for raw in body.splitlines():
                for part in simpleSplit(raw, "Times-Roman", 8, X(x2 - x1 - 12)) or [""]:
                    if yy > 226:
                        break
                    c.drawString(X(x1 + 3), Y(yy), part)
                    yy += 10.5
        return 300

    def table_frame(top, last):
        head_bottom = top + 24
        box(TX1, top, TX2, head_bottom, fill=0.9)
        c.setFont("Helvetica", 8.5)
        for label, x1, x2 in zip(("Item", "Description", "Qty", "Rate", "Amount"), COLS, COLS[1:]):
            c.drawCentredString(X((x1 + x2) / 2), Y(top + 15), label)
        for x in COLS[1:-1]:
            vline(x, top, BODY_BOTTOM)
        vline(TX1, top, BODY_BOTTOM)
        vline(TX2, top, BODY_BOTTOM)
        hline(TX1, TX2, BODY_BOTTOM)
        if last:
            box(TX1, BODY_BOTTOM, TX2, TOTAL_BOTTOM)
            vline(367, BODY_BOTTOM, TOTAL_BOTTOM)
            c.setFont("Helvetica-Bold", 13)
            c.drawString(X(380), Y(718), "Total")
            c.setFont("Times-Roman", 8.5)
            c.drawRightString(X(551), Y(718), f"${total:,.2f}")
        return head_bottom

    LH = 10.5  # row line height

    # Work out which rows go on which page first, so the Total is only drawn on the last page.
    desc_w = X(COLS[2] - COLS[1] - 8)

    def wrap(line):
        return simpleSplit(str(line["description"]), "Times-Roman", 8, desc_w) or [""]

    pages: list[list[dict]] = [[]]
    first_top = 300 + 24
    cont_top = 70 + 24
    used = 0.0
    limit = BODY_BOTTOM - first_top - 8
    for line in lines:
        wrapped = wrap(line)
        h = max(len(wrapped), 1) * LH + 3
        if used + h > limit and pages[-1]:
            pages.append([])
            used = 0.0
            limit = BODY_BOTTOM - cont_top - 8
        pages[-1].append((line, wrapped))
        used += h

    for n, rows in enumerate(pages):
        last = n == len(pages) - 1
        if n == 0:
            top = page_one_header()
        else:
            top = 70
            c.setFont("Helvetica", 9)
            c.drawString(X(36), Y(52), f"Purchase Order {po['po_number']} (continued, page {n + 1})")
        y = table_frame(top, last) + 14
        for line, wrapped in rows:
            c.setFont("Times-Roman", 8)
            c.drawString(X(COLS[0] + 3), Y(y), str(line.get("item") or ""))
            for k, part in enumerate(wrapped):
                c.drawString(X(COLS[1] + 4), Y(y + k * LH), part)
            c.drawRightString(X(COLS[3] - 4), Y(y), str(int(line["quantity"])))
            c.drawRightString(X(COLS[4] - 4), Y(y), f"{_money(line['rate']):,.2f}")
            c.drawRightString(X(COLS[5] - 4), Y(y), f"{line_amount(line):,.2f}")
            y += max(len(wrapped), 1) * LH + 3
        if last and po.get("note"):
            c.setFont("Times-Italic", 8)
            c.drawString(X(COLS[1] + 4), Y(y + 6), f"Note: {po['note']}"[:110])
        c.showPage()
    c.save()
    return buf.getvalue()


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
def build_xlsx(po: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"PO {po['po_number']}"
    bold = Font(bold=True)
    thin = Side(style="thin", color="000000")
    grey = PatternFill("solid", fgColor="E6E6E6")

    for n, t in enumerate(COMPANY_LINES):
        ws.cell(row=1 + n, column=1, value=t).font = Font(bold=(n == 0))
    ws["D1"] = "Purchase Order"
    ws["D1"].font = Font(bold=True, size=16)
    ws["D2"], ws["E2"] = "Date", _date_text(po.get("date"))
    ws["D3"], ws["E3"] = "P.O. No.", po["po_number"]
    for ref in ("D2", "D3"):
        ws[ref].font = bold

    ws["A5"], ws["D5"] = "Vendor", "Ship To"
    ws["A5"].font = ws["D5"].font = bold
    vendor = [po["vendor_name"]] + [l for l in (po.get("vendor_address") or "").splitlines() if l.strip()]
    ship = [l for l in (po.get("ship_to") or "").splitlines() if l.strip()]
    for n in range(max(len(vendor), len(ship))):
        if n < len(vendor):
            ws.cell(row=6 + n, column=1, value=vendor[n])
        if n < len(ship):
            ws.cell(row=6 + n, column=4, value=ship[n])

    head = 6 + max(len(vendor), len(ship), 4) + 1
    for col, label in enumerate(("Item", "Description", "Qty", "Rate", "Amount"), start=1):
        cell = ws.cell(row=head, column=col, value=label)
        cell.font, cell.fill = bold, grey
        cell.alignment = Alignment(horizontal="center")
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    r = head + 1
    for line in po["lines"]:
        ws.cell(row=r, column=1, value=line.get("item") or "")
        ws.cell(row=r, column=2, value=line["description"]).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=3, value=int(line["quantity"]))
        ws.cell(row=r, column=4, value=float(_money(line["rate"]))).number_format = "#,##0.00"
        ws.cell(row=r, column=5, value=float(line_amount(line))).number_format = "#,##0.00"
        r += 1
    ws.cell(row=r + 1, column=4, value="Total").font = Font(bold=True, size=12)
    tot = ws.cell(row=r + 1, column=5, value=float(po_total(po["lines"])))
    tot.font, tot.number_format = Font(bold=True, size=12), '"$"#,##0.00'
    if po.get("note"):
        ws.cell(row=r + 3, column=1, value=f"Note: {po['note']}")

    for col, width in zip("ABCDE", (18, 60, 8, 12, 14)):
        ws.column_dimensions[col].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
